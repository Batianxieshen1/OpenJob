# -*- coding: utf-8 -*-
"""用户自维护素材表 → 标准素材库格式转换器（一次性适配，不改动原文件）"""
from pathlib import Path

from openpyxl import Workbook, load_workbook

SRC = Path(r'C:\Users\暴龙战士wink\Desktop\个人简历\简历素材库.xlsx')
OUT = Path(r'C:\Users\暴龙战士wink\Desktop\agent\OpenJob\data\resume_materials.xlsx')

HEADERS = [
    "id", "type", "title", "organization", "role", "start_date", "end_date",
    "description", "achievements", "skills", "keywords", "target_directions",
    "source", "resume_allowed", "priority", "notes",
]


def cell(row, header_map, name):
    idx = header_map.get(name)
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    return str(v).strip() if v is not None else ""


def norm_date(raw):
    """2025.04 / 2025.9 / 2025 → 2025-04 / 2025-09 / 2025；至今 原样。"""
    raw = (raw or "").strip()
    if not raw or raw == "至今":
        return raw
    parts = raw.replace("-", ".").split(".")
    if len(parts) == 1 and parts[0].isdigit():
        return parts[0]
    if len(parts) >= 2 and parts[0].isdigit():
        year = parts[0]
        month = parts[1].zfill(2) if parts[1].isdigit() else ""
        if month:
            return f"{year}-{month}"
        return year
    return raw


def stars_to_priority(raw):
    count = (raw or "").count("⭐")
    return max(1, min(5, count)) if count else 3


def classify_experience(category):
    if "实习" in category:
        return "experience"
    if "项目" in category:
        return "project"
    if "校园" in category or "实践" in category:
        return "campus_activity"
    if "技能" in category:
        return "skill_evidence"
    return "experience"


def main():
    wb = load_workbook(SRC, data_only=True, read_only=True)
    materials = []

    # ---------- 📋 经历 ----------
    sheet = wb["📋 经历"]
    rows = list(sheet.iter_rows(values_only=True))
    hm = {h.strip(): i for i, h in enumerate(rows[0]) if h}
    for row in rows[1:]:
        if not cell(row, hm, "名称 / 标题"):
            continue
        seq = cell(row, hm, "序号") or len(materials) + 1
        category = cell(row, hm, "类别")
        description = cell(row, hm, "详细描述（STAR 法则")
        achievements = cell(row, hm, "量化成果")
        if not description:
            description = f"{cell(row, hm, '名称 / 标题')}：{cell(row, hm, '角色 / 职位') or '参与'}，{achievements or '详见成果'}"
        skills = cell(row, hm, "关键技能")
        target = cell(row, hm, "适用岗位")
        materials.append({
            "id": f"exp_{seq}",
            "type": classify_experience(category),
            "title": cell(row, hm, "名称 / 标题"),
            "organization": cell(row, hm, "组织 / 机构"),
            "role": cell(row, hm, "角色 / 职位"),
            "start_date": norm_date(cell(row, hm, "开始时间")),
            "end_date": norm_date(cell(row, hm, "结束时间")),
            "description": description,
            "achievements": achievements,
            "skills": skills,
            "keywords": "、".join(x for x in (skills, target) if x),
            "target_directions": target,
            "source": cell(row, hm, "证明材料"),
            "resume_allowed": "是",
            "priority": stars_to_priority(cell(row, hm, "重要程度")),
            "notes": (f"原类别：{category}；{cell(row, hm, '备注')}")[:500],
        })

    # ---------- 🏆 奖项 ----------
    sheet = wb["🏆 奖项"]
    rows = list(sheet.iter_rows(values_only=True))
    hm = {h.strip(): i for i, h in enumerate(rows[0]) if h}
    for row in rows[1:]:
        title = cell(row, hm, "奖项全称")
        if not title:
            continue
        seq = cell(row, hm, "序号") or len(materials) + 1
        level = cell(row, hm, "级别")
        rank = cell(row, hm, "排名/等级")
        ratio = cell(row, hm, "获奖比例")
        description = cell(row, hm, "描述 / 参赛作品")
        if not description:
            description = f"{title}（{level or '奖项'}）"
        achievements = "；".join(x for x in (level, f"排名/等级：{rank}" if rank else "", f"获奖比例：{ratio}" if ratio else "") if x)
        materials.append({
            "id": f"awd_{seq}",
            "type": "award",
            "title": title,
            "organization": cell(row, hm, "颁发机构"),
            "role": "",
            "start_date": norm_date(cell(row, hm, "获奖时间")),
            "end_date": "",
            "description": description,
            "achievements": achievements,
            "skills": "",
            "keywords": level,
            "target_directions": cell(row, hm, "适用岗位"),
            "source": cell(row, hm, "证明材料"),
            "resume_allowed": "是",
            "priority": stars_to_priority(cell(row, hm, "重要程度")),
            "notes": cell(row, hm, "备注")[:500],
        })

    # ---------- 📜 证明 → certification ----------
    sheet = wb["📜 证明"]
    rows = list(sheet.iter_rows(values_only=True))
    hm = {h.strip(): i for i, h in enumerate(rows[0]) if h}
    for row in rows[1:]:
        title = cell(row, hm, "证书全称")
        if not title:
            continue
        seq = cell(row, hm, "序号") or len(materials) + 1
        category = cell(row, hm, "类别")
        score = cell(row, hm, "分数 / 等级")
        cert_no = cell(row, hm, "证书编号")
        description = f"证书类别：{category or '证书'}"
        if score:
            description += f"；分数/等级：{score}"
        if cert_no:
            description += f"；证书编号：{cert_no}"
        materials.append({
            "id": f"cert_{seq}",
            "type": "certification",
            "title": title,
            "organization": cell(row, hm, "颁发机构"),
            "role": "",
            "start_date": norm_date(cell(row, hm, "获取时间")),
            "end_date": norm_date(cell(row, hm, "有效期至")),
            "description": description,
            "achievements": score,
            "skills": "",
            "keywords": category,
            "target_directions": cell(row, hm, "适用岗位"),
            "source": cell(row, hm, "证明材料路径"),
            "resume_allowed": "是",
            "priority": stars_to_priority(cell(row, hm, "重要程度")),
            "notes": cell(row, hm, "备注")[:500],
        })

    wb.close()

    # ---------- 生成标准 XLSX ----------
    out = Workbook()
    sheet = out.active
    sheet.title = "素材库"
    sheet.append(HEADERS)
    for m in materials:
        sheet.append([
            m["id"], m["type"], m["title"], m["organization"], m["role"],
            m["start_date"], m["end_date"], m["description"], m["achievements"],
            m["skills"], m["keywords"], m["target_directions"], m["source"],
            m["resume_allowed"], m["priority"], m["notes"],
        ])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    out.save(OUT)
    print(f"✓ 转换完成：{len(materials)} 条素材 → {OUT}")
    print("  类型分布:", {t: sum(1 for m in materials if m['type'] == t) for t in set(m['type'] for m in materials)})


if __name__ == "__main__":
    from openpyxl import Workbook  # noqa: F401
    main()
