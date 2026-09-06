"""本地简历素材库：XLSX 模板、解析、校验、哈希与原子保存。

素材库是用户自行维护的 XLSX 源文件（默认 data/resume_materials.xlsx），
上传成功后解析为带 SHA-256 的 JSON 索引；简历生成只读取已校验的索引。
`source`、`notes` 和 `resume_allowed=否` 的素材绝不进入 AI prompt。
"""

from __future__ import annotations

import hashlib
import json
import os
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

MAX_XLSX_BYTES = 10 * 1024 * 1024

DEFAULT_HEADERS = [
    "id", "type", "title", "organization", "role", "start_date", "end_date",
    "description", "achievements", "skills", "keywords", "target_directions",
    "source", "resume_allowed", "priority", "notes",
]

SHEET_NAME = "素材库"

MATERIAL_TYPES = {
    "experience", "project", "award", "student_work",
    "campus_activity", "skill_evidence", "certification", "other",
}

_TEXT_LIMITS = {
    "id": (1, 64),
    "title": (1, 120),
    "description": (1, 2000),
    "achievements": (0, 2000),
    "source": (0, 500),
    "notes": (0, 500),
    "organization": (0, 120),
    "role": (0, 80),
}

_ID_PATTERN = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-")


class MaterialLibraryError(ValueError):
    """素材库文件无效。"""


@dataclass(frozen=True)
class MaterialLibrary:
    items: list[dict]
    source_name: str
    sha256: str
    updated_at: str

    @property
    def count(self) -> int:
        return len(self.items)


def split_cell_list(value: object) -> list[str]:
    """拆分逗号/中文逗号/顿号/斜杠分隔的单元格为去空白列表。"""
    import re

    raw = str(value or "").strip()
    if not raw:
        return []
    parts = re.split(r"[,，、/]+", raw)
    return [p.strip() for p in parts if p.strip()]


def _cell(value) -> str:
    """公式单元格只读缓存值；None 按空字符串处理。"""
    if value is None:
        return ""
    return str(value).strip()


def _validate_id(raw: str, row_no: int) -> str:
    if not raw:
        raise MaterialLibraryError(f"第 {row_no} 行：id 不能为空")
    if not (1 <= len(raw) <= 64):
        raise MaterialLibraryError(f"第 {row_no} 行：id 长度必须为 1-64")
    if any(ch not in _ID_PATTERN for ch in raw):
        raise MaterialLibraryError(f"第 {row_no} 行：id 只允许字母、数字、下划线和连字符")
    return raw


def _validate_type(raw: str, row_no: int) -> str:
    if raw not in MATERIAL_TYPES:
        raise MaterialLibraryError(
            f"第 {row_no} 行：type 非法（{raw or '空'}），只允许 {'/'.join(sorted(MATERIAL_TYPES))}"
        )
    return raw


def _validate_date(raw: str, row_no: int, field: str) -> str:
    if not raw or raw == "至今":
        if field == "start_date" and raw == "至今":
            raise MaterialLibraryError(f"第 {row_no} 行：start_date 不允许为 至今")
        return raw
    for fmt in ("%Y", "%Y-%m"):
        try:
            datetime.strptime(raw, fmt)
            return raw
        except ValueError:
            continue
    raise MaterialLibraryError(f"第 {row_no} 行：{field} 格式非法（{raw}），只允许 YYYY 或 YYYY-MM")


def _validate_bool(raw: str, row_no: int) -> bool:
    lowered = raw.lower()
    if raw in {"是", "true", "1"} or lowered == "true":
        return True
    if raw in {"否", "false", "0"} or lowered == "false":
        return False
    raise MaterialLibraryError(f"第 {row_no} 行：resume_allowed 只允许 是/否、true/false、1/0")


def _validate_priority(raw: str, row_no: int) -> int:
    if not raw:
        return 3
    try:
        value = int(raw)
    except ValueError as exc:
        raise MaterialLibraryError(f"第 {row_no} 行：priority 必须是 1-5 的整数") from exc
    if not 1 <= value <= 5:
        raise MaterialLibraryError(f"第 {row_no} 行：priority 必须是 1-5 的整数")
    return value


def parse_workbook(content: bytes, *, filename: str) -> MaterialLibrary:
    """解析 XLSX 字节为素材库；任何校验失败都让整次上传失败。"""
    if not filename.lower().endswith(".xlsx"):
        raise MaterialLibraryError("只支持 .xlsx 格式（不接受 .xls 或宏文件）")
    if len(content) > MAX_XLSX_BYTES:
        raise MaterialLibraryError("文件超过 10MB 上限")
    if not content.startswith(b"PK"):
        raise MaterialLibraryError("文件不是有效的 XLSX（ZIP）格式")

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise MaterialLibraryError(f"XLSX 无法解析：{exc}") from exc

    if SHEET_NAME not in workbook.sheetnames:
        raise MaterialLibraryError(f"缺少名为「{SHEET_NAME}」的工作表")
    sheet = workbook[SHEET_NAME]

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        raise MaterialLibraryError("工作表为空：第一行必须是列名")

    headers = [_cell(h) for h in header_row]
    missing = [h for h in DEFAULT_HEADERS if h not in headers]
    if missing:
        raise MaterialLibraryError(f"缺少必需列：{'、'.join(missing)}")
    index_of = {name: headers.index(name) for name in DEFAULT_HEADERS}

    items: list[dict] = []
    seen_ids: set[str] = set()
    row_no = 1
    for row in rows_iter:
        row_no += 1
        cells = list(row or [])
        if not any(_cell(c) for c in cells):
            continue  # 空行跳过，但不能用它绕过“至少一条素材”

        item: dict = {}
        item["id"] = _validate_id(_cell(cells[index_of["id"]]) if index_of["id"] < len(cells) else "", row_no)
        if item["id"] in seen_ids:
            raise MaterialLibraryError(f"第 {row_no} 行：重复 id「{item['id']}」")
        seen_ids.add(item["id"])

        item["type"] = _validate_type(_cell(cells[index_of["type"]]) if index_of["type"] < len(cells) else "", row_no)

        for field in ("title", "description"):
            raw = _cell(cells[index_of[field]]) if index_of[field] < len(cells) else ""
            low, high = _TEXT_LIMITS[field]
            if not (low <= len(raw) <= high):
                raise MaterialLibraryError(f"第 {row_no} 行：{field} 长度必须为 {low}-{high} 字符")
            item[field] = raw

        for field in ("organization", "role", "source", "notes"):
            raw = _cell(cells[index_of[field]]) if index_of[field] < len(cells) else ""
            low, high = _TEXT_LIMITS[field]
            if len(raw) > high:
                raise MaterialLibraryError(f"第 {row_no} 行：{field} 超过 {high} 字符上限")
            item[field] = raw

        item["start_date"] = _validate_date(
            _cell(cells[index_of["start_date"]]) if index_of["start_date"] < len(cells) else "",
            row_no, "start_date",
        )
        end_raw = _cell(cells[index_of["end_date"]]) if index_of["end_date"] < len(cells) else ""
        item["end_date"] = _validate_date(end_raw, row_no, "end_date")
        if (
            item["start_date"] and item["end_date"] and item["end_date"] != "至今"
            and item["end_date"] < item["start_date"]
        ):
            raise MaterialLibraryError(f"第 {row_no} 行：end_date 不得早于 start_date")

        item["achievements"] = _cell(cells[index_of["achievements"]]) if index_of["achievements"] < len(cells) else ""
        if len(item["achievements"]) > 2000:
            raise MaterialLibraryError(f"第 {row_no} 行：achievements 超过 2000 字符上限")

        item["skills"] = split_cell_list(cells[index_of["skills"]] if index_of["skills"] < len(cells) else None)
        item["keywords"] = split_cell_list(cells[index_of["keywords"]] if index_of["keywords"] < len(cells) else None)
        item["target_directions"] = split_cell_list(
            cells[index_of["target_directions"]] if index_of["target_directions"] < len(cells) else None
        )

        allowed_raw = _cell(cells[index_of["resume_allowed"]]) if index_of["resume_allowed"] < len(cells) else ""
        item["resume_allowed"] = _validate_bool(allowed_raw, row_no)
        item["priority"] = _validate_priority(
            _cell(cells[index_of["priority"]]) if index_of["priority"] < len(cells) else "", row_no
        )

        items.append(item)

    if not items:
        raise MaterialLibraryError("至少需要一条素材数据行（空行不算）")

    workbook.close()
    return MaterialLibrary(
        items=items,
        source_name=filename,
        sha256=hashlib.sha256(content).hexdigest(),
        updated_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def build_template_workbook() -> bytes:
    """生成可直接填写的素材库模板（含表头与一行示例说明）。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(DEFAULT_HEADERS)
    sheet.append([
        "exp_001",
        "experience",
        "示例：某项目/实习名称",
        "示例公司或组织",
        "实习生",
        "2024-03",
        "2024-08",
        "示例：做了什么、怎么做的（必填，1-2000 字）",
        "示例：量化结果，如 阅读量提升 30%",
        "技能A, 技能B",
        "用于匹配 JD 的关键词",
        "数据分析, 运营",
        "来源备注（不会进入简历和 AI）",
        "是",
        3,
        "私有备注（不会进入简历和 AI）",
    ])
    buf = BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def load_index(index_path: Path) -> MaterialLibrary:
    """读取已校验的 JSON 索引；损坏时抛 MaterialLibraryError。"""
    try:
        data = json.loads(index_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise MaterialLibraryError("素材库索引不存在，请重新上传 XLSX") from exc
    except json.JSONDecodeError as exc:
        raise MaterialLibraryError("素材库索引损坏，请重新上传 XLSX 或刷新索引") from exc

    items = data.get("items")
    if not isinstance(items, list) or not items:
        raise MaterialLibraryError("素材库索引为空或损坏")
    return MaterialLibrary(
        items=items,
        source_name=str(data.get("source_name") or "resume_materials.xlsx"),
        sha256=str(data.get("sha256") or ""),
        updated_at=str(data.get("updated_at") or ""),
    )


def save_library_atomically(
    library: MaterialLibrary,
    content: bytes,
    *,
    xlsx_path: Path,
    index_path: Path,
) -> None:
    """XLSX 与 JSON 索引原子落盘：先写 .tmp + fsync，再 os.replace。"""
    for path in (xlsx_path, index_path):
        path.parent.mkdir(parents=True, exist_ok=True)

    index_payload = json.dumps(
        {
            "source_name": library.source_name,
            "sha256": library.sha256,
            "updated_at": library.updated_at,
            "items": library.items,
        },
        ensure_ascii=False,
        indent=1,
    )

    for target, payload in ((xlsx_path, content), (index_path, index_payload.encode("utf-8"))):
        tmp = target.with_suffix(target.suffix + ".tmp")
        with open(tmp, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp, target)


def resolve_library_paths(config: dict, data_dir: Path) -> tuple[Path, Path]:
    """解析并校验素材 XLSX 与索引路径，二者必须位于 data_dir 内。

    - 从 config["profile"]["resume_materials_path"] 读取，未配置用默认；
    - 相对路径按 data_dir.parent（项目根目录）解析，不按进程工作目录；
    - 越界路径抛 MaterialLibraryError，不静默切换文件。
    """
    raw = str((config.get("profile") or {}).get("resume_materials_path") or "./data/resume_materials.xlsx")
    data_root = Path(data_dir).resolve()
    candidate = Path(raw)
    if not candidate.is_absolute():
        candidate = data_root.parent / candidate
    candidate = candidate.resolve()
    try:
        candidate.relative_to(data_root)
    except ValueError as exc:
        raise MaterialLibraryError("简历素材库路径必须位于项目 data 目录内") from exc
    index_path = candidate.parent / (candidate.stem + ".index.json")
    return candidate, index_path


def load_or_refresh_library(materials_path: Path, index_path: Path) -> MaterialLibrary:
    """生成期入口：文件存在但索引缺失/哈希不一致/索引损坏时重新解析当前 XLSX。"""
    if not materials_path.exists():
        raise MaterialLibraryError("未找到简历素材库，请在配置页上传 .xlsx")
    content = materials_path.read_bytes()
    sha = hashlib.sha256(content).hexdigest()

    if index_path.exists():
        try:
            cached = load_index(index_path)
            if cached.sha256 == sha:
                return cached
        except MaterialLibraryError:
            pass  # 索引损坏/为空：按缺失处理，重新解析

    library = parse_workbook(content, filename=materials_path.name)
    save_library_atomically(library, content, xlsx_path=materials_path, index_path=index_path)
    return library
