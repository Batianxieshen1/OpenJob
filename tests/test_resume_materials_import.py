"""智能导入领域层测试：扫描/映射/归一化/示例排除/稳定 ID/隐私边界。"""

import io
import json
from pathlib import Path

import io
import pytest
import unittest
from openpyxl import Workbook

from openjob.resume_materials_import import (
    analyze_workbook,
    commit_import,
    infer_column_mapping,
    infer_header_row,
    normalize_import_rows,
)


def build_sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def make_excel(sheets: dict[str, tuple[list[str], list[dict]]], *, leading_rows: int = 0, hidden: list[str] | None = None):
    wb = Workbook()
    wb.remove(wb.active)
    hidden = hidden or []
    first = True
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        for _ in range(leading_rows):
            ws.append(["说明文字行", None, None])
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        if name in hidden:
            ws.sheet_state = "hidden"
    # openpyxl 不允许全部工作表隐藏：补一个可见占位表
    if all(ws.sheet_state != "visible" for ws in wb.worksheets):
        wb.create_sheet("占位")
    first = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CHINESE_HEADERS = ["名称 / 标题", "组织 / 机构", "角色 / 职位", "开始时间", "结束时间",
                   "详细描述（STAR 法则", "量化成果", "关键技能", "适用岗位"]


def _cn_row(**overrides):
    row = {
        "名称 / 标题": "赛事运营实习", "组织 / 机构": "省篮协", "角色 / 职位": "志愿者统筹",
        "开始时间": "2025.04", "结束时间": "2025.06",
        "详细描述（STAR 法则": "统筹 200 名志愿者保障赛事",
        "量化成果": "9 条检票通道零差错", "关键技能": "团队管理",
        "适用岗位": "运营,数据分析",
    }
    row.update(overrides)
    return row


class TestHeaderInference:
    def test_header_on_first_row(self):
        wb = Workbook()
        ws = wb.active
        ws.append(CHINESE_HEADERS)
        ws.append(["实习", "公司", "角色", "2024.01", "2024.06", "描述", "成果", "技能", "岗位"])
        assert infer_header_row([list(r) for r in ws.iter_rows(values_only=True)]) == 0

    def test_header_after_title_and_blank_rows(self):
        rows = [
            ["我的简历素材", None, None, None, None, None, None, None, None],
            [None] * 9,
            ["更新于 2024 年", None, None, None, None, None, None, None, None],
            [None] * 9,
            list(CHINESE_HEADERS),
            ["实习", "公司", "角色", "2024.01", "2024.06", "描述", "成果", "技能", "岗位"],
        ]
        assert infer_header_row(rows) == 4

    def test_unreliable_header_returns_none(self):
        rows = [["随便一行文字"], ["另一行"], ["再来一行"]]
        assert infer_header_row(rows) is None


class TestColumnMapping:
    def test_chinese_headers_map_to_standard_fields(self):
        mapping = infer_column_mapping(CHINESE_HEADERS, [[]])
        by_field = {m["target_field"]: m for m in mapping}
        assert by_field["title"]["source_column"] == "名称 / 标题"
        assert by_field["description"]["source_column"] == "详细描述（STAR 法则"
        assert by_field["achievements"]["source_column"] == "量化成果"
        assert by_field["target_directions"]["source_column"] == "适用岗位"

    def test_english_headers_map(self):
        mapping = infer_column_mapping(
            ["title", "description", "organization", "skills"], [[]]
        )
        by_field = {m["target_field"]: m for m in mapping}
        assert by_field["title"]["source_column"] == "title"

    def test_unknown_column_becomes_ignore(self):
        mapping = infer_column_mapping(["乱写的列", "标题"], [["x", "y"]])
        by_source = {m["source_column"]: m for m in mapping}
        assert by_source["乱写的列"]["target_field"] == "__ignore__"
        assert by_source["标题"]["target_field"] == "title"


class TestAnalyzeWorkbook:
    def test_analyze_multisheet_chinese_workbook(self):
        content = make_excel({
            "项目经历": (["项目名称", "项目内容", "成果"], [
                {"项目名称": "电商分析框架", "项目内容": "SQL 与 Python", "成果": "开源"},
            ]),
            "奖项": (["奖项全称", "获奖时间", "级别"], [
                {"奖项全称": "市级奖", "获奖时间": "2025.12", "级别": "市级"},
            ]),
        })
        session = analyze_workbook(content, filename="我的.xlsx")
        assert session.source_sha256
        names = [s["name"] for s in session.sheets]
        assert "项目经历" in names and "奖项" in names
        proj = next(s for s in session.sheets if s["name"] == "项目经历")
        assert proj["inferred_type"] == "project"
        assert proj["row_count"] == 1

    def test_analyze_does_not_write_official_library(self, tmp_path):
        content = make_excel({"项目经历": (["项目名称", "项目内容"], [{"项目名称": "P", "项目内容": "C"}])})
        analyze_workbook(content, filename="a.xlsx", imports_dir=tmp_path)
        assert not (tmp_path / "resume_materials.xlsx").exists()

    def test_hidden_sheet_flagged(self):
        content = make_excel(
            {"项目经历": (["项目名称", "项目内容"], [{"项目名称": "P", "项目内容": "C"}])},
            hidden=["项目经历"],
        )
        session = analyze_workbook(content, filename="a.xlsx")
        assert session.sheets[0]["hidden"] is True
        assert any("隐藏" in w for w in session.warnings)

    def test_chinese_and_space_path(self, tmp_path):
        src = tmp_path / "我的 素材表.xlsx"
        src.write_bytes(make_excel({
            "项目经历": (["项目名称", "项目内容"], [{"项目名称": "P", "项目内容": "C"}]),
        }))
        session = analyze_workbook(src.read_bytes(), filename=src.name)
        assert session.filename == src.name

    def test_oversize_rejected(self):
        content = b"PK" + b"x" * (11 * 1024 * 1024)
        with pytest.raises(Exception, match="10MB"):
            analyze_workbook(content, filename="big.xlsx")

    def test_non_xlsx_rejected(self):
        with pytest.raises(Exception, match="xlsx"):
            analyze_workbook(b"not excel", filename="a.xls")


class TestNormalizeAndCommit:
    def _confirmed(self, tmp_path, sheets, confirm_extra=None):
        content = make_excel(sheets)
        session = analyze_workbook(content, filename="我的.xlsx", imports_dir=tmp_path)
        confirm = {
            "session": session.to_storage_dict(),
            "sheets": [
                {
                    "name": s["name"],
                    "include": True,
                    "header_row": s["header_row"],
                    "field_mapping": {
                        m["source_column"]: m["target_field"]
                        for m in s["mapping"]
                        if m["target_field"] != "__ignore__"
                    },
                    "type_override": s.get("inferred_type") or "other",
                    "excluded_rows": [],
                }
                for s in session.sheets
            ],
            "defaults": {"resume_allowed": True, "priority": 3},
        }
        if confirm_extra:
            for sheet_conf in confirm["sheets"]:
                sheet_conf.update(confirm_extra.get(sheet_conf["name"], {}))
            confirm.update({k: v for k, v in confirm_extra.items() if k in ("defaults",)})
        return content, session, confirm

    def test_normalize_maps_chinese_columns_and_generates_stable_ids(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "项目内容", "成果"],
                [{"项目名称": "电商框架", "项目内容": "SQL 与 Python", "成果": "开源"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        items, warnings = normalize_import_rows(content, confirm)
        assert len(items) == 1
        assert items[0]["title"] == "电商框架"
        assert "SQL" in items[0]["description"] or "SQL" in items[0]["achievements"]
        assert items[0]["id"].startswith("imp_")
        # 稳定 ID：同内容重新归一化，ID 不变
        items2, _ = normalize_import_rows(content, confirm)
        assert items2[0]["id"] == items[0]["id"]

    def test_example_rows_are_excluded(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "项目内容"],
                [
                    {"项目名称": "__openjob_example__", "项目内容": "示例数据，请替换"},
                    {"项目名称": "真实项目", "项目内容": "真实描述"},
                ],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        items, warnings = normalize_import_rows(content, confirm)
        assert [i["title"] for i in items] == ["真实项目"]
        assert any("示例" in w for w in warnings)

    def test_type_column_overrides_sheet_inference(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "项目内容", "类别"],
                [{"项目名称": "校园志愿", "项目内容": "d", "类别": "校园实践"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets, {"项目经历": {"type_override": "campus_activity"}})
        items, _ = normalize_import_rows(content, confirm)
        assert items[0]["type"] == "campus_activity"

    def test_date_formats_normalized(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "项目内容", "开始时间", "结束时间"],
                [{"项目名称": "P", "项目内容": "C", "开始时间": "2024.9", "结束时间": "至今"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        items, _ = normalize_import_rows(content, confirm)
        assert items[0]["start_date"] == "2024-09"
        assert items[0]["end_date"] == "至今"

    def test_skills_split_by_semicolon_and_newline_not_slash(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "项目内容", "关键技能"],
                [{"项目名称": "P", "项目内容": "C",
                  "关键技能": "SQL; Python\n数据分析/BI，机器学习"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        items, _ = normalize_import_rows(content, confirm)
        # 技能按分号/换行/逗号/顿号拆，不按斜杠拆
        assert "数据分析/BI" in items[0]["skills"]

    def test_rows_missing_description_excluded_with_warning(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "成果"],
                [{"项目名称": "坏行-缺描述", "成果": "有成果"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        # 映射里没有 description 列 → 全部行缺描述 → 排除
        items, warnings = normalize_import_rows(content, confirm)
        assert items == []
        assert any("缺少事实描述" in w for w in warnings)

    def test_commit_strict_and_atomic(self, tmp_path):
        sheets = {
            "项目经历": (
                ["项目名称", "项目内容", "成果"],
                [{"项目名称": "电商框架", "项目内容": "SQL 与 Python", "成果": "开源"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        items, warnings = normalize_import_rows(content, confirm)
        xlsx_path = tmp_path / "official.xlsx"
        index_path = tmp_path / "official.index.json"
        library = commit_import(content, items, xlsx_path=xlsx_path, index_path=index_path)
        assert library.count == 1
        assert index_path.exists()
        # 原 XLSX 与索引均落盘
        assert xlsx_path.exists()

    def test_commit_missing_description_fails_strict(self, tmp_path):
        # normalize 排除缺描述行后 items 为空 → commit 报“至少需要一条”
        sheets = {
            "项目经历": (
                ["项目名称"],
                [{"项目名称": "坏行"}],
            ),
        }
        content, session, confirm = self._confirmed(tmp_path, sheets)
        items, _ = normalize_import_rows(content, confirm)
        from openjob.resume_materials import MaterialLibraryError

        with pytest.raises(MaterialLibraryError, match="至少需要一条"):
            commit_import(content, items, xlsx_path=tmp_path / "o.xlsx", index_path=tmp_path / "o.json")


class ResourceLimitTests(unittest.TestCase):
    """Task 1：250 行完整导入、超限拒绝、不创建会话/不覆盖正式库。"""

    def setUp(self):
        import tempfile

        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _build_250(self):
        rows = [
            {"项目名称": f"项目{i:03d}", "项目内容": f"描述{i}：SQL 与 Python 数据处理",
             "成果": f"成果{i}", "类别": "项目"}
            for i in range(250)
        ]
        content = make_excel({"项目经历": (["项目名称", "项目内容", "成果"], rows)})
        return content

    def test_analyze_and_normalize_all_250_rows_without_truncation(self):
        from openjob.resume_materials_import import analyze_workbook, normalize_import_rows

        content = self._build_250()
        session = analyze_workbook(content, filename="250.xlsx")
        sheet = session.sheets[0]
        self.assertEqual(sheet["row_count"], 250)

        confirm = {
            "session": session.to_storage_dict(),
            "sheets": [{
                "name": "项目经历", "include": True, "header_row": sheet["header_row"],
                "field_mapping": {m["source_column"]: m["target_field"] for m in sheet["mapping"]},
                "type_override": "project", "excluded_rows": [],
            }],
            "defaults": {"resume_allowed": True, "priority": 3},
        }
        items, _ = normalize_import_rows(content, confirm)
        self.assertEqual(len(items), 250)
        self.assertEqual(items[0]["title"], "项目000")
        self.assertEqual(items[-1]["title"], "项目249")

        from openjob.resume_materials_import import commit_import
        xlsx = self.tmp / "official.xlsx"
        index = self.tmp / "official.index.json"
        library = commit_import(content, items, xlsx_path=xlsx, index_path=index)
        self.assertEqual(library.count, 250)

    def test_over_row_limit_is_rejected_instead_of_partially_imported(self):
        from openjob.resume_materials import MaterialLibraryError
        from openjob.resume_materials_import import analyze_workbook

        rows = [
            {"项目名称": f"项目{i}", "项目内容": f"描述{i}", "成果": f"成果{i}", "类别": "项目"}
            for i in range(5001)
        ]
        content = make_excel({"项目经历": (["项目名称", "项目内容", "成果"], rows)})
        with self.assertRaises(MaterialLibraryError) as ctx:
            analyze_workbook(content, filename="5001.xlsx")
        message = str(ctx.exception)
        self.assertIn("5001", message)
        self.assertIn("5000", message)

    def test_over_column_limit_is_rejected(self):
        from openjob.resume_materials import MaterialLibraryError
        from openjob.resume_materials_import import analyze_workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "项目经历"
        ws.append([f"列{i}" for i in range(101)])
        ws.append(["值" for _ in range(101)])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(MaterialLibraryError) as ctx:
            analyze_workbook(buf.getvalue(), filename="101cols.xlsx")
        self.assertIn("101", str(ctx.exception))
        self.assertIn("100", str(ctx.exception))

    def test_over_limit_analysis_does_not_create_session_or_replace_library(self):
        from openjob.resume_materials import MaterialLibraryError
        from openjob.resume_materials_import import analyze_workbook

        imports_dir = self.tmp / "imports"
        imports_dir.mkdir()
        before = imports_dir.exists()

        # 先建立一份正式库（走 normalize+commit 的标准路径）
        from openjob.resume_materials_import import commit_import

        official_xlsx = self.tmp / "official.xlsx"
        official_index = self.tmp / "official.index.json"
        base_content = make_excel({"项目经历": (["项目名称", "项目内容", "成果"], [{"项目名称": "P", "项目内容": "C"}])})
        session = analyze_workbook(base_content, filename="base.xlsx")
        items, _ = normalize_import_rows(base_content, {
            "session": session.to_storage_dict(),
            "sheets": [{"name": s["name"], "include": True,
                        "header_row": s["header_row"],
                        "field_mapping": {m["source_column"]: m["target_field"] for m in s["mapping"]},
                        "type_override": s.get("inferred_type") or "other", "excluded_rows": []}
                       for s in session.sheets],
            "defaults": {"resume_allowed": True, "priority": 3},
        })
        commit_import(base_content, items, xlsx_path=official_xlsx, index_path=official_index)
        xlsx_before = official_xlsx.read_bytes()
        index_before = official_index.read_bytes()

        rows = [
            {"项目名称": f"项目{i}", "项目内容": f"描述{i}", "成果": f"成果{i}", "类别": "项目"}
            for i in range(5001)
        ]
        over = make_excel({"项目经历": (["项目名称", "项目内容", "成果"], rows)})
        with self.assertRaises(MaterialLibraryError):
            analyze_workbook(over, filename="over.xlsx", imports_dir=imports_dir)

        sessions = [d for d in imports_dir.iterdir() if d.is_dir()] if imports_dir.exists() else []
        self.assertEqual(len(sessions), 0)
        self.assertEqual(official_xlsx.read_bytes(), xlsx_before)
        self.assertEqual(official_index.read_bytes(), index_before)


class SessionCleanupTests(unittest.TestCase):
    """TTL 清理：只删过期合法会话，不动哨兵/越界目录/有效会话。"""

    def setUp(self):
        import tempfile

        self._tmp = tempfile.TemporaryDirectory()
        self.imports_dir = Path(self._tmp.name) / "resume_material_imports"
        self.imports_dir.mkdir(parents=True)
        self.sentinel = Path(self._tmp.name) / "哨兵.txt"
        self.sentinel.write_text("keep", encoding="utf-8")
        self.outside = Path(self._tmp.name) / "outside-dir"
        self.outside.mkdir()

    def tearDown(self):
        self._tmp.cleanup()

    def test_cleanup_removes_only_expired_sessions_inside_import_root(self):
        from datetime import datetime, timedelta, timezone

        from openjob.resume_materials_import import cleanup_expired_import_sessions

        expired = self.imports_dir / ("a" * 32)
        expired.mkdir()
        created = datetime.now(timezone.utc) - timedelta(hours=25)
        (expired / "analysis.json").write_text(json.dumps({
            "import_id": "a" * 32, "created_at": created.isoformat(),
        }), encoding="utf-8")

        removed = cleanup_expired_import_sessions(self.imports_dir)
        self.assertGreaterEqual(removed, 1)
        self.assertFalse(expired.exists())
        self.assertTrue(self.sentinel.exists())
        # 哨兵文件（imports 根外）必须仍然存在——清理不越界
        self.assertTrue((self.imports_dir.parent / self.sentinel.name).exists())

    def test_cleanup_keeps_active_sessions(self):
        from datetime import datetime, timedelta, timezone

        from openjob.resume_materials_import import cleanup_expired_import_sessions

        # 23h59m 前创建的会话应保留
        active = self.imports_dir / ("d" * 32)
        active.mkdir()
        created = datetime.now(timezone.utc) - timedelta(hours=23, minutes=59)
        (active / "analysis.json").write_text(json.dumps({
            "import_id": "d" * 32, "created_at": created.isoformat(),
        }), encoding="utf-8")
        removed = cleanup_expired_import_sessions(self.imports_dir)
        self.assertTrue(active.exists())
        self.assertEqual(removed, 0)
