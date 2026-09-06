"""转换脚本 CLI 测试：dry-run/force/缺表/重复 id/中文路径/输出可回解析。"""

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "convert_user_materials.py"
ROOT = SCRIPT.parents[1]


def build_user_source(path: Path, rows_by_sheet: dict[str, list[dict]] | None = None):
    """构造与用户自维护格式一致的源文件（可只含部分表以测试缺表场景）。"""
    all_headers = {
        "📋 经历": ["序号", "开始时间", "结束时间", "类别", "名称 / 标题", "组织 / 机构", "城市",
                    "角色 / 职位", "详细描述（STAR 法则", "量化成果", "关键技能", "证明材料",
                    "重要程度", "简历状态", "适用岗位", "备注"],
        "🏆 奖项": ["序号", "获奖时间", "类别", "奖项全称", "颁发机构", "级别", "排名/等级",
                   "获奖比例", "描述 / 参赛作品", "证明材料", "重要程度", "简历状态", "适用岗位", "备注"],
        "📜 证明": ["序号", "获取时间", "有效期至", "类别", "证书全称", "颁发机构", "分数 / 等级",
                   "证书编号", "证明材料路径", "重要程度", "简历状态", "适用岗位", "备注"],
    }
    wb = Workbook()
    wb.remove(wb.active)
    sheets = rows_by_sheet or {
        "📋 经历": [
            {"序号": 1, "开始时间": "2025.04", "结束时间": "2025.06", "类别": "实习经历",
             "名称 / 标题": "赛事运营实习", "组织 / 机构": "省篮协", "城市": "佛山",
             "角色 / 职位": "志愿者统筹", "详细描述（STAR 法则": "统筹 200 名志愿者",
             "量化成果": "9 条通道零差错", "关键技能": "团队管理", "证明材料": "",
             "重要程度": "⭐⭐⭐⭐", "简历状态": "已使用", "适用岗位": "运营,数据分析", "备注": ""},
        ],
        "🏆 奖项": [
            {"序号": 1, "获奖时间": "2025.12", "奖项全称": "市级实践奖", "颁发机构": "梅州市团委",
             "级别": "市级", "描述 / 参赛作品": "调研作品获奖", "重要程度": "⭐⭐⭐", "备注": ""},
        ],
        "📜 证明": [
            {"序号": 1, "获取时间": "2025.9", "证书全称": "CET-4", "类别": "语言",
             "分数 / 等级": "474", "证明材料路径": "", "备注": ""},
        ],
    }
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(all_headers[sheet_name])
        for row in rows:
            ws.append([row.get(h, "") for h in all_headers[sheet_name]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def run_cli(*args):
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )


class ConverterCliTests(unittest.TestCase):
    def setUp(self):
        import tempfile

        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.src = self.tmp / "我的 素材表.xlsx"  # 中文 + 空格路径
        build_user_source(self.src)

    def tearDown(self):
        self._tmp.cleanup()

    def test_dry_run_parses_without_writing(self):
        dst = self.tmp / "out.xlsx"
        result = run_cli("--input", str(self.src), "--output", str(dst), "--dry-run")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("3 条素材", result.stdout)
        self.assertIn("dry-run", result.stdout)
        self.assertFalse(dst.exists())

    def test_default_refuses_overwrite_force_overwrites(self):
        dst = self.tmp / "out.xlsx"
        dst.write_bytes(b"old")
        result = run_cli("--input", str(self.src), "--output", str(dst))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("目标已存在", result.stdout + result.stderr)

        result = run_cli("--input", str(self.src), "--output", str(dst), "--force")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(dst.exists())

    def test_output_reparseable_by_parse_workbook(self):
        dst = self.tmp / "out.xlsx"
        result = run_cli("--input", str(self.src), "--output", str(dst), "--force")
        self.assertEqual(result.returncode, 0)

        sys_path = str(ROOT / "src")
        if sys_path not in sys.path:
            sys.path.insert(0, sys_path)
        from openjob.resume_materials import parse_workbook

        library = parse_workbook(dst.read_bytes(), filename="out.xlsx")
        self.assertEqual(library.count, 3)
        type_ids = {m["type"] for m in library.items}
        self.assertIn("experience", type_ids)
        self.assertIn("award", type_ids)
        self.assertIn("certification", type_ids)

    def test_missing_sheet_fails(self):
        src = self.tmp / "缺表.xlsx"
        build_user_source(src, rows_by_sheet={"📋 经历": [
            {"序号": 1, "开始时间": "2025.04", "类别": "实习经历",
             "名称 / 标题": "实习", "详细描述（STAR 法则": "d", "resume_allowed": "是"},
        ]})
        result = run_cli("--input", str(src), "--output", str(self.tmp / "o.xlsx"))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("缺少工作表", result.stdout + result.stderr)

    def test_duplicate_id_fails(self):
        src = self.tmp / "重复.xlsx"
        build_user_source(src, rows_by_sheet={
            "📋 经历": [
                {"序号": 1, "类别": "实习经历", "名称 / 标题": "A",
                 "详细描述（STAR 法则": "d", "resume_allowed": "是"},
                {"序号": 1, "类别": "项目", "名称 / 标题": "B",
                 "详细描述（STAR 法则": "d", "resume_allowed": "是"},
            ],
            "🏆 奖项": [], "📜 证明": [],
        })
        result = run_cli("--input", str(src), "--output", str(self.tmp / "o.xlsx"))
        self.assertNotEqual(result.returncode, 0)

    def test_chinese_space_path_ok(self):
        dst = self.tmp / "输出 目标.xlsx"
        result = run_cli("--input", str(self.src), "--output", str(dst), "--force")
        self.assertEqual(result.returncode, 0)
        self.assertTrue(dst.exists())


if __name__ == "__main__":
    unittest.main()
